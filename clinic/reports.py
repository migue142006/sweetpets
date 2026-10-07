from collections import defaultdict
from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape
from django.db.models import Sum
from django.http import HttpResponse
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from .models import InvoiceItem, Payment, Service, ZERO

def report_data(start,end,kind=''):
    items=InvoiceItem.objects.filter(invoice__issued_at__date__gte=start,invoice__issued_at__date__lte=end).select_related('invoice','service')
    if kind: items=items.filter(service__kind=kind)
    revenue,cost=ZERO,ZERO
    services=defaultdict(lambda:{'events':0,'pets':set(),'total':ZERO})
    for i in items:
        revenue+=i.total-i.tax; cost+=i.unit_cost*i.quantity
        key=i.service.get_kind_display()
        services[key]['events']+=i.quantity
        if i.invoice.pet_id: services[key]['pets'].add(i.invoice.pet_id)
        services[key]['total']+=i.total
    payments=Payment.objects.filter(created_at__date__gte=start,created_at__date__lte=end).select_related('invoice')
    months=defaultdict(lambda:ZERO)
    # Pagos mixtos se distribuyen proporcionalmente entre líneas para filtrar por servicio.
    for p in payments:
        amount=p.amount
        if kind:
            total=p.invoice.total
            selected=sum((i.total for i in p.invoice.items.filter(service__kind=kind)),ZERO)
            amount=(amount*selected/total).quantize(Decimal('.01')) if total else ZERO
        months[p.created_at.strftime('%Y-%m')]+=amount
    count=sum(v['events'] for v in services.values())
    rows=[[k,v['events'],len(v['pets']),round(v['events']*100/count,2) if count else 0,v['total']] for k,v in sorted(services.items())]
    return {'rows':rows,'months':dict(sorted(months.items())),'income':sum(months.values(),ZERO),'net_sales':revenue,'direct_cost':cost,'margin':revenue-cost,'events':count}

def pdf_response(filename,title,paragraphs,rows):
    buffer=BytesIO(); styles=getSampleStyleSheet()
    story=[Paragraph(escape(title),styles['Title']),Spacer(1,12)]
    for text in paragraphs: story.extend([Paragraph(escape(str(text)),styles['BodyText']),Spacer(1,8)])
    if rows:
        cells=[[Paragraph(escape(str(v)),styles['BodyText']) for v in row] for row in rows]
        table=Table(cells,repeatRows=1,hAlign='LEFT')
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e8f3ef')),('GRID',(0,0),(-1,-1),.4,colors.grey),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
        story.append(table)
    SimpleDocTemplate(buffer).build(story)
    response=HttpResponse(buffer.getvalue(),content_type='application/pdf')
    response['Content-Disposition']=f'attachment; filename="{filename}"'
    return response

def excel_response(data):
    book=Workbook(); sheet=book.active; sheet.title='Servicios'
    sheet.append(['Servicio','Eventos o unidades','Mascotas distintas','Porcentaje de eventos','Facturado con impuestos'])
    for row in data['rows']: sheet.append(row)
    sheet.freeze_panes='A2'
    for column in ['A','B','C','D','E']: sheet.column_dimensions[column].width=27
    monthly=book.create_sheet('Recaudo mensual'); monthly.append(['Mes','Recaudo'])
    for month,amount in data['months'].items(): monthly.append([month,amount])
    totals=book.create_sheet('Totales')
    for key in ['income','net_sales','direct_cost','margin','events']: totals.append([key,data[key]])
    stream=BytesIO(); book.save(stream)
    r=HttpResponse(stream.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    r['Content-Disposition']='attachment; filename="sweetpets-reporte.xlsx"'
    return r
